"""
main.py — FastAPI application entry point.

Run with:
    uvicorn app.main:app --reload --port 8000

Authentication
--------------
All /api/* routes (except /api/auth/login) require a JWT in the header:
    Authorization: Bearer <token>

Role-based access is enforced via FastAPI dependencies at the router level.
Access to data is deny-by-default (see docs/PERFORMANCE_GOVERNANCE.md):
  - Organisation-wide dashboards/reports/exports → require_org_wide (a grant on the root unit, or admin)
  - /api/platform/* and module APIs → scope resolved per request; results limited to granted units
  - /api/position/*           → limited to the units in the user's scope
  - /api/upload/*             → require_admin     (admin only)
  - /api/records/export/csv   → require_export    (admin or user; not viewer)
  - /api/admin/*              → require_admin     (admin only)
  - /api/integration/*        → require_admin     (admin only)
  - /api/ingest/{source}      → per-source push token (X-Madzi-Ingest-Token), no user session
"""
from contextlib import asynccontextmanager

from fastapi import Depends, FastAPI, Request
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse, HTMLResponse, JSONResponse, Response
from fastapi.staticfiles import StaticFiles
import os
import time
import uuid

from app.auth import ensure_default_admin, get_current_user, require_admin
from app.core.config import settings
from app.core.logging import REQUEST_ID_CTX, logger as app_logger
from app.database import SessionLocal, create_tables
from app.migrate import SchemaNotReady
from app.routers import analytics, benchmarking, budget, catalogue, compliance, fiscal_years, integration, panels, records, report_generator, reports, strategic, upload, insights
from app.routers import config as config_router
from app.routers.users import admin_router, auth_router
from app import model_registry as _models  # noqa: F401  (every table registered before start-up)
from app.platform.errors import PlatformError
from app.platform.router import router as platform_router
from app.platform.scope import require_org_wide
from app.core.limiter import limiter as _limiter
from slowapi import _rate_limit_exceeded_handler
from slowapi.errors import RateLimitExceeded
from slowapi.middleware import SlowAPIMiddleware


@asynccontextmanager
async def lifespan(app: FastAPI):
    """Startup: create DB tables, bootstrap default admin if needed."""
    settings.validate_startup()
    try:
        create_tables()
    except SchemaNotReady as exc:
        print(f"[STOP] {exc}")
        raise
    db = SessionLocal()
    try:
        ensure_default_admin(db)
        from app.platform.bootstrap import at_startup as _platform_startup
        _platform_startup(db)
        # Auto-import: if records table is empty and an Excel file exists, seed it
        from app.database import Record
        count = db.query(Record).count()
        if count == 0:
            _auto_import(db)
    finally:
        db.close()
    print("[OK] Database tables ready")
    print("[OK] User authentication active  (JWT / bcrypt)")
    yield


def _auto_import(db):
    """Seed DB from RawData.xlsx if available. Called only when records table is empty."""
    import glob
    base = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    candidates = (
        glob.glob(os.path.join(base, "uploads", "RawData*.xlsx"))
        + glob.glob(os.path.join(base, "data", "RawData*.xlsx"))
    )
    if not candidates:
        print("[WARN] Records table is empty. Upload a RawData workbook from Administration > Upload,")
        print("   or place RawData*.xlsx in uploads/ or data/ and restart.")
        return

    xlsx_path = candidates[0]
    print(f"⚙  Auto-importing from {xlsx_path} ...")
    try:
        import openpyxl
        from app.database import Record
        from app.services.excel_parser import ExcelParser, resolve_column

        wb = openpyxl.load_workbook(xlsx_path, data_only=True)
        ws = wb["DataEntry"] if "DataEntry" in wb.sheetnames else wb.active

        # Detect header row (row 1 or 2)
        row1 = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
        row2 = [c.value for c in next(ws.iter_rows(min_row=2, max_row=2))]
        if any(str(h).strip() in ("Zone", "Scheme", "Month", "Year") for h in (row2 or [])):
            headers = [str(c).strip() if c else "" for c in row2]
            data_start = 3
        else:
            headers = [str(c).strip() if c else "" for c in row1]
            data_start = 2

        # Header → Record column mapping
        HMAP = {
            "Zone": "zone", "Scheme": "scheme", "Fiscal Year": "fiscal_year",
            "Year": "year", "Month No.": "month_no", "Month": "month", "Quarter": "quarter",
            "Volume Produced (m³)": "vol_produced",
            "Vol Billed Individual Postpaid": "vol_billed_indiv_pp",
            "Vol Billed CWP Postpaid": "vol_billed_cwp_pp",
            "Vol Billed Institutions Postpaid": "vol_billed_inst_pp",
            "Vol Billed Commercial Postpaid": "vol_billed_comm_pp",
            "TOTAL Vol Billed Postpaid": "total_vol_billed_pp",
            "Vol Billed Individual Prepaid": "vol_billed_indiv_prepaid",
            "Vol Billed CWP Prepaid": "vol_billed_cwp_prepaid",
            "Vol Billed Institutions Prepaid": "vol_billed_inst_prepaid",
            "Vol Billed Commercial Prepaid": "vol_billed_comm_prepaid",
            "TOTAL Vol Billed Prepaid": "total_vol_billed_prepaid",
            "TOTAL Revenue Water m³": "revenue_water", "Non-Revenue Water m³": "nrw",
            "% NRW": "pct_nrw",
            "Chlorine kg": "chlorine_kg", "Alum Sulphate kg": "alum_kg",
            "Soda Ash kg": "soda_ash_kg", "Algae Floc litres": "algae_floc_litres",
            "Sud Floc litres": "sud_floc_litres", "Potassium Permanganate kg": "kmno4_kg",
            "Cost of Chemicals": "chem_cost", "Chem Cost per m³": "chem_cost_per_m3",
            "Power Usage kWh": "power_kwh", "Cost of Power": "power_cost",
            "Power Cost per m³": "power_cost_per_m3",
            "Distances Covered km": "distances_km", "Fuel Used litres": "fuel_used_litres",
            "Cost of Fuel": "fuel_cost", "Maintenance": "maintenance",
            "Staff Costs": "staff_costs", "Wages": "wages",
            "Other Overhead": "other_overhead",
            "TOTAL Operating Costs": "op_cost",
            "OpCost per m³ Produced": "op_cost_per_m3_produced",
            "OpCost per m³ Billed": "op_cost_per_m3_billed",
            "Permanent Staff": "perm_staff", "Temporary Staff": "temp_staff",
            "ALL Conn BroughtFwd": "all_conn_bfwd", "ALL Conn Applied": "all_conn_applied",
            "ALL Conn TOTAL Done": "new_connections", "ALL Conn CarriedFwd": "all_conn_cfwd",
            "Prepaid Meters Installed": "prepaid_meters_installed",
            "Disconnected Individual": "disconnected_individual",
            "Disconnected Institutional": "disconnected_inst",
            "Disconnected Commercial": "disconnected_commercial",
            "Disconnected CWP": "disconnected_cwp", "TOTAL Disconnected": "total_disconnected",
            "Active Postpaid Individual": "active_post_individual",
            "Active Postpaid Institutional": "active_post_inst",
            "Active Postpaid Commercial": "active_post_commercial",
            "Active Postpaid CWP": "active_post_cwp",
            "TOTAL Active Postpaid": "active_postpaid",
            "Active Prepaid Individual": "active_prep_individual",
            "Active Prepaid Institutional": "active_prep_inst",
            "Active Prepaid Commercial": "active_prep_commercial",
            "Active Prepaid CWP": "active_prep_cwp",
            "TOTAL Active Prepaid": "active_prepaid",
            "TOTAL Active Customers": "active_customers",
            "Total Metered Consumers": "total_metered",
            "Population Supply Area": "pop_supply_area",
            "Population Supplied": "pop_supplied",
            "Pct Population Supplied": "pct_pop_supplied",
            "ALL StuckM BroughtFwd": "stuck_meters", "ALL StuckM New": "stuck_new",
            "ALL StuckM Repaired": "stuck_repaired", "ALL StuckM Replaced": "stuck_replaced",
            "TOTAL Pipe Breakdowns": "pipe_breakdowns", "Pump Breakdowns": "pump_breakdowns",
            "Pump Hours Lost": "pump_hours_lost",
            "Normal Supply Hours": "supply_hours", "Power Failure Hours": "power_fail_hours",
            "DevLines 32mm": "dev_lines_32mm", "DevLines 50mm": "dev_lines_50mm",
            "DevLines 63mm": "dev_lines_63mm", "DevLines 90mm": "dev_lines_90mm",
            "DevLines 110mm": "dev_lines_110mm", "TOTAL Dev Lines Done": "dev_lines_total",
            "TOTAL Cash Coll PP": "cash_coll_pp", "TOTAL Cash Coll Prepaid": "cash_coll_prepaid",
            "TOTAL Cash Collected": "cash_collected",
            "TOTAL Amt Billed PP": "amt_billed_pp",
            "TOTAL Amt Billed Prepaid": "amt_billed_prepaid",
            "TOTAL Amount Billed": "amt_billed",
            "TOTAL Service Charge": "service_charge", "TOTAL Meter Rental": "meter_rental",
            "TOTAL Sales": "total_sales",
            "Private Debtors": "private_debtors", "Public Debtors": "public_debtors",
            "TOTAL Debtors": "total_debtors",
            "OpCost per Sales": "op_cost_per_sales",
            "Cash Collection Rate": "collection_rate",
            "Collection per Total Sales": "collection_per_sales",
            "Cust Applied Connection": "conn_applied",
            "Days to Quotation": "days_to_quotation",
            "Cust Fully Paid": "conn_fully_paid", "Days to Connect": "days_to_connect",
            "Connectivity Rate": "connectivity_rate",
            "Queries Received": "queries_received",
            "Time to Resolve Queries": "time_to_resolve",
            "Response Time avg": "response_time_avg",
        }

        rec_cols = {c.name for c in Record.__table__.columns if c.name != "id"}
        str_cols = {"zone", "scheme", "fiscal_year", "month", "quarter"}

        col_map = {}
        for i, h in enumerate(headers):
            parser_col = resolve_column(ExcelParser._normalize_header(h))
            if parser_col in rec_cols:
                col_map[i] = parser_col
            elif h in HMAP and HMAP[h] in rec_cols:
                col_map[i] = HMAP[h]

        inserted = 0
        for row in ws.iter_rows(min_row=data_start, max_row=ws.max_row, values_only=True):
            data = {}
            for i, db_col in col_map.items():
                v = row[i] if i < len(row) else None
                if db_col in str_cols:
                    data[db_col] = str(v) if v is not None else ""
                else:
                    try:
                        data[db_col] = float(v) if v is not None else 0.0
                    except (ValueError, TypeError):
                        data[db_col] = 0.0
            if not data.get("zone") or data["zone"] in ("Zone", ""):
                continue
            data["year"] = int(data.get("year", 0))
            data["month_no"] = int(data.get("month_no", 0))
            db.add(Record(**data))
            inserted += 1

        db.commit()
        print(f"[OK] Auto-imported {inserted} records from {os.path.basename(xlsx_path)}")
    except Exception as e:
        print(f"[FAIL] Auto-import failed: {e}")
        db.rollback()


app = FastAPI(
    title="MadziHub API",
    description=(
        "MadziHub — water utility performance & intelligence platform.\n\n"
        "Monetary values are in the installation's reporting currency and the "
        "financial year follows its configured start month (see `GET /api/config`). "
        "Volume in **m³**.\n\n"
        "**Authentication:** `POST /api/auth/login` with username + password "
        "to obtain a Bearer token.  Include it as:\n"
        "`Authorization: Bearer <token>`\n\n"
        "**Roles:** `admin` · `user` · `viewer`"
    ),
    version="2.0.0",
    contact={"name": "MadziHub"},
    license_info={"name": "Internal Use"},
    lifespan=lifespan,
)

app.state.limiter = _limiter
app.add_exception_handler(RateLimitExceeded, _rate_limit_exceeded_handler)


@app.exception_handler(PlatformError)
async def _platform_error(_request: Request, exc: PlatformError):
    return JSONResponse(status_code=exc.status_code, content={"detail": exc.detail})
app.add_middleware(SlowAPIMiddleware)

# ── CORS ──────────────────────────────────────────────────────
# Set MADZI_ALLOWED_ORIGINS to a comma-separated list of origins in
# production, e.g. "https://hub.example-utility.org"
# Restricted localhost defaults for development; production validation blocks '*'.
_allowed_origins = settings.allowed_origins

app.add_middleware(
    CORSMiddleware,
    allow_origins=_allowed_origins,
    allow_credentials=False,
    allow_methods=["GET", "POST", "PUT", "DELETE"],
    allow_headers=["*"],
)

# ── Auth endpoints (public — no auth dependency) ───────────────
app.include_router(auth_router)

# ── Installation config (public branding + authenticated client config) ──
app.include_router(config_router.router)

# ── Admin user-management (admin role required) ───────────────
app.include_router(admin_router, dependencies=[Depends(require_admin)])

# ── Organisation-wide data (any authenticated user with an organisation-wide grant) ──
# These dashboards, reports and exports aggregate every region. A user whose access
# is limited to particular units gets 403 here and works in the scoped modules instead.
# require_export on /export/csv is enforced inside records.py
_org_wide = [Depends(require_org_wide)]
app.include_router(records.router,       dependencies=_org_wide)
app.include_router(analytics.router,    dependencies=_org_wide)
app.include_router(budget.router,       dependencies=_org_wide)
app.include_router(catalogue.router,    dependencies=_org_wide)
app.include_router(fiscal_years.router, dependencies=_org_wide)
app.include_router(panels.router,       dependencies=_org_wide)
app.include_router(compliance.router,   dependencies=_org_wide)
app.include_router(benchmarking.router, dependencies=_org_wide)
app.include_router(reports.router,          dependencies=_org_wide)
app.include_router(report_generator.router, dependencies=_org_wide)
app.include_router(insights.router,     dependencies=_org_wide)
app.include_router(strategic.router,    dependencies=_org_wide)

# ── Shared governance foundation and modules (scope resolved per request) ──
from app.modules.strategy.router import router as strategy_router  # noqa: E402
from app.modules.scorecard.router import router as scorecard_router  # noqa: E402

app.include_router(platform_router)
app.include_router(strategy_router)
app.include_router(scorecard_router)

# ── Upload (admin only) ───────────────────────────────────────
app.include_router(upload.router, dependencies=[Depends(require_admin)])

# Integration hub: sources and catalogue (admin), strategic position (users),
# push ingest (per-source machine token, checked in the handler).
app.include_router(integration.admin_router,    dependencies=[Depends(require_admin)])
app.include_router(integration.position_router, dependencies=[Depends(get_current_user)])
app.include_router(integration.ingest_router)

# ── Static assets ─────────────────────────────────────────────
STATIC_DIR = os.path.join(os.path.dirname(__file__), "static")
INDEX_PATH  = os.path.join(STATIC_DIR, "index.html")
APP_CORE_JS = os.path.join(STATIC_DIR, "assets", "js", "app-core.js")


def _brand_values() -> dict:
    """Tenant identity with admin org-profile overrides applied."""
    from app.core.tenant import tenant
    from app.database import OrgProfile
    ident = tenant.identity
    db = SessionLocal()
    try:
        profile = db.query(OrgProfile).filter(OrgProfile.id == 1).first()
    finally:
        db.close()
    name = (profile and profile.org_name) or ident.name
    return {
        "product_title": ident.product_title,
        "tagline": ident.tagline,
        # "product": no tenant logo, so the page shows the full MadziHub lockup.
        # "tenant": the utility's own logo, title and tagline.
        "brand_mode": "tenant" if tenant.logo_path else "product",
        # Organisation line under the sidebar wordmark, blank when it would repeat the product name.
        "org_line": "" if name.strip().lower() == ident.product_title.strip().lower() else name,
        "name": name,
        "short_name": (profile and profile.short_name) or ident.short_name,
        "currency": tenant.currency.code,
        "currency_symbol": tenant.currency.symbol,
        "plan_title": tenant.strategic_plan.title,
        "country": (profile and profile.country) or ident.country or "",
        "zone_count": str(len(tenant.hierarchy.zones)),
        "zone_plural": tenant.hierarchy.levels[0].plural if tenant.hierarchy.levels else "Zones",
    }


def _js_text(value: str) -> str:
    """Make a value safe inside any JS string literal ('', "", ``) and in innerHTML.

    "$" is escaped rather than dropped: currency symbols such as "$" and "US$" must
    survive, and \\u0024 yields "$" without ever opening a ${...} substitution.
    """
    import json
    cleaned = "".join(ch for ch in (value or "") if ch not in "<>\"'`\\")
    return json.dumps(cleaned)[1:-1].replace("$", "\\u0024")


_JS_CACHE: dict = {}


# Registered before the /static mount so it takes precedence over the static file.
@app.get("/static/assets/js/app-core.js", include_in_schema=False)
def serve_app_core_js(request: Request):
    """Serve app-core.js with tenant placeholders filled (labels, currency, targets, calendar)."""
    import hashlib
    import json
    from app.core.tenant import tenant
    from app.utils import MONTHS_ORDER

    brand = _brand_values()
    subs = {
        "__ORG_SHORT__": _js_text(brand["short_name"]),
        "__ORG_NAME_COUNTRY__": _js_text(" · ".join(v for v in (brand["name"], brand["country"]) if v)),
        "__ORG_NAME__": _js_text(brand["name"]),
        "__PRINT_LOGO__": ("/static/brand/madzihub-logo-light.svg" if brand["brand_mode"] == "product"
                           else "/api/config/logo"),
        "__PLAN_TITLE__": _js_text(brand["plan_title"]),
        "__CURRENCY__": _js_text(brand["currency"]),
        "__CUR_SYM__": _js_text(brand["currency_symbol"]),
        "__NRW_TARGET__": f"{tenant.target('nrw_pct', 25.0):g}",
        "__FY_MONTHS__": json.dumps(MONTHS_ORDER),
        "__ZONE_COLORS__": json.dumps(tenant.zone_colors),
    }
    key = (os.path.getmtime(APP_CORE_JS), tuple(sorted(subs.items())))
    if key not in _JS_CACHE:
        _JS_CACHE.clear()
        with open(APP_CORE_JS, encoding="utf-8") as f:
            content = f.read()
        for placeholder, value in subs.items():
            content = content.replace(placeholder, value)
        _JS_CACHE[key] = (content, hashlib.sha256(content.encode()).hexdigest()[:32])
    content, etag = _JS_CACHE[key]
    headers = {"Cache-Control": "no-cache", "ETag": f'"{etag}"'}
    if request.headers.get("if-none-match") == f'"{etag}"':
        return Response(status_code=304, headers=headers)
    return Response(content, media_type="application/javascript", headers=headers)
if os.path.isdir(STATIC_DIR):
    app.mount("/static", StaticFiles(directory=STATIC_DIR), name="static")


@app.middleware("http")
async def add_request_context(request: Request, call_next):
    request_id = request.headers.get("X-Request-ID") or str(uuid.uuid4())
    token = REQUEST_ID_CTX.set(request_id)
    start = time.perf_counter()
    try:
        response = await call_next(request)
    except Exception:
        duration_ms = round((time.perf_counter() - start) * 1000, 2)
        app_logger.exception(
            "request_failed",
            extra={
                "method": request.method,
                "path": request.url.path,
                "duration_ms": duration_ms,
                "request_id": request_id,
            },
        )
        REQUEST_ID_CTX.reset(token)
        raise
    duration_ms = round((time.perf_counter() - start) * 1000, 2)
    response.headers["X-Request-ID"] = request_id
    app_logger.info(
        "request_completed",
        extra={
            "method": request.method,
            "path": request.url.path,
            "status_code": response.status_code,
            "duration_ms": duration_ms,
            "request_id": request_id,
        },
    )
    REQUEST_ID_CTX.reset(token)
    return response


# ── Root — inject API base URL and tenant branding, then serve dashboard ──
def _brand_html(content: str) -> str:
    """Fill the tenant placeholders in index.html (HTML-escaped)."""
    from html import escape
    brand = _brand_values()
    values = {
        "__PRODUCT_TITLE__": brand["product_title"],
        "__TAGLINE__": brand["tagline"],
        "__BRAND_MODE__": brand["brand_mode"],
        "__ORG_LINE__": brand["org_line"],
        "__ORG_NAME__": brand["name"],
        "__ORG_SHORT__": brand["short_name"],
        "__CURRENCY__": brand["currency"],
        "__PLAN_TITLE__": brand["plan_title"],
        "__ORG_COUNTRY__": brand["country"],
        "__ZONE_COUNT__": brand["zone_count"],
        "__ZONE_PLURAL__": brand["zone_plural"],
    }
    for key, value in values.items():
        content = content.replace(key, escape(value or ""))
    return content



@app.get("/", include_in_schema=False)
async def serve_dashboard(request: Request):
    if not os.path.exists(INDEX_PATH):
        return {"message": "MadziHub API running. Place index.html in app/static/"}
    base_url = str(request.base_url).rstrip("/")
    with open(INDEX_PATH, encoding="utf-8") as f:
        content = f.read()
    content = content.replace("__API_BASE__", base_url)
    content = _brand_html(content)
    return HTMLResponse(
        content=content,
        headers={
            "Cache-Control": "no-store, no-cache, must-revalidate, max-age=0",
            "Pragma": "no-cache",
        }
    )


# Browsers (and /docs) request /favicon.ico at the root regardless of <link> tags.
@app.get("/favicon.ico", include_in_schema=False)
def favicon():
    return FileResponse(os.path.join(STATIC_DIR, "brand", "favicon.ico"),
                        headers={"Cache-Control": "public, max-age=86400"})


# ── Health check (public) ─────────────────────────────────────
@app.get("/health", tags=["System"])
def health():
    return {"status": "ok", "version": app.version}


@app.get("/api/debug/db-status", tags=["System"], dependencies=[Depends(require_admin)])
def db_status():
    """Admin-only diagnostic endpoint — shows record counts per year/month."""
    from app.database import Record
    db = SessionLocal()
    try:
        total = db.query(Record).count()
        from sqlalchemy import func
        breakdown = (
            db.query(Record.year, Record.month_no, func.count())
            .group_by(Record.year, Record.month_no)
            .order_by(Record.year, Record.month_no)
            .all()
        )
        sample = db.query(Record).first()
        sample_data = {}
        if sample:
            for col in ("zone", "scheme", "year", "month_no", "month",
                        "vol_produced", "amt_billed", "cash_collected", "op_cost",
                        "active_customers", "total_metered", "stuck_meters", "nrw"):
                sample_data[col] = getattr(sample, col, None)
        return {
            "total_records": total,
            "by_year_month": [{"year": y, "month_no": m, "count": c} for y, m, c in breakdown],
            "sample_row": sample_data,
        }
    finally:
        db.close()
