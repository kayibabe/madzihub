"""
Upload checks and text extraction.

check(): the extension must be allowed for the document type AND the bytes must actually be
that kind of file (magic numbers / container structure), so a renamed executable or script is
refused. Size is limited by UPLOAD_LIMIT_MB.

extract(): text for search. Failure is recorded on the version and shown; the stored file is
never changed. OCR is not included (scanned PDFs extract as "empty").
"""
from __future__ import annotations

import io
import re
import zipfile

from app.core.config import settings
from app.platform.errors import Invalid

MIME = {"pdf": "application/pdf", "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "csv": "text/csv",
        "txt": "text/plain", "md": "text/markdown", "png": "image/png", "jpg": "image/jpeg", "jpeg": "image/jpeg"}
TEXT_TYPES = ("csv", "txt", "md")
MAX_TEXT = 2_000_000   # characters kept for search


def extension(filename: str) -> str:
    name = (filename or "").strip().lower()
    return name.rsplit(".", 1)[-1] if "." in name else ""


def _looks_like(ext: str, data: bytes) -> bool:
    if ext == "pdf":
        return data[:5] == b"%PDF-"
    if ext in ("docx", "xlsx"):
        if data[:4] != b"PK\x03\x04":
            return False
        try:
            names = set(zipfile.ZipFile(io.BytesIO(data)).namelist())
        except zipfile.BadZipFile:
            return False
        part = "word/document.xml" if ext == "docx" else "xl/workbook.xml"
        return "[Content_Types].xml" in names and part in names
    if ext == "png":
        return data[:8] == b"\x89PNG\r\n\x1a\n"
    if ext in ("jpg", "jpeg"):
        return data[:3] == b"\xff\xd8\xff"
    if ext in TEXT_TYPES:
        if b"\x00" in data[:65536]:
            return False
        try:
            data.decode("utf-8")
        except UnicodeDecodeError:
            try:
                data.decode("cp1252")
            except UnicodeDecodeError:
                return False
        head = data[:512].lstrip().lower()
        return not head.startswith((b"<script", b"<!doctype html", b"<html", b"mz"))
    return False


def check(filename: str, data: bytes, allowed: list[str]) -> tuple[str, str]:
    """Validate an upload; returns (extension, mime)."""
    ext = extension(filename)
    if not data:
        raise Invalid("The file is empty.")
    if len(data) > settings.upload_limit_mb * 1024 * 1024:
        raise Invalid(f"The file is larger than {settings.upload_limit_mb} MB.")
    if ext not in MIME or ext not in [a.lower() for a in allowed]:
        raise Invalid(f"'.{ext or '?'}' files are not accepted for this document type "
                      f"(allowed: {', '.join('.' + a for a in allowed)}).")
    if data[:2] == b"MZ" or data[:4] == b"\x7fELF":
        raise Invalid("Executable files are not accepted.")
    if not _looks_like(ext, data):
        raise Invalid(f"The file's content does not match its '.{ext}' extension, so it was refused.")
    return ext, MIME[ext]


def _xml_text(xml: bytes) -> str:
    text = re.sub(rb"</w:p>|</w:tr>", b"\n", xml)
    text = re.sub(rb"<[^>]+>", b" ", text)
    import html
    return html.unescape(re.sub(r"[ \t]+", " ", text.decode("utf-8", "ignore")))


def extract(ext: str, data: bytes) -> tuple[str | None, str, str | None]:
    """(text, status, error). Status: ok | empty | unsupported | failed."""
    try:
        if ext in TEXT_TYPES:
            try:
                text = data.decode("utf-8")
            except UnicodeDecodeError:
                text = data.decode("cp1252")
        elif ext == "docx":
            with zipfile.ZipFile(io.BytesIO(data)) as z:
                text = _xml_text(z.read("word/document.xml"))
        elif ext == "xlsx":
            from openpyxl import load_workbook
            wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            parts = []
            for ws in wb.worksheets:
                parts.append(ws.title)
                for row in ws.iter_rows(values_only=True):
                    cells = [str(v) for v in row if v is not None]
                    if cells:
                        parts.append(" ".join(cells))
            wb.close()
            text = "\n".join(parts)
        elif ext == "pdf":
            from pypdf import PdfReader
            reader = PdfReader(io.BytesIO(data))
            text = "\n".join((page.extract_text() or "") for page in reader.pages)
        else:
            return None, "unsupported", None
    except Exception as exc:   # extraction must never block or alter the upload
        return None, "failed", f"{type(exc).__name__}: {str(exc)[:200]}"
    text = (text or "").strip()
    if not text:
        return None, "empty", "No text found (a scanned document needs OCR, which is not included)."
    return text[:MAX_TEXT], "ok", None
